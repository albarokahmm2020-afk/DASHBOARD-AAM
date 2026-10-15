"""Monthly Meeting Kasie deck generator based on the supplied 2026 template."""

from io import BytesIO
from pathlib import Path

import pandas as pd

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(68, 114, 196)
PALE_BLUE = RGBColor(221, 235, 247)
LIGHT = RGBColor(242, 246, 250)
GREEN = RGBColor(198, 239, 206)
RED = RGBColor(255, 220, 220)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(35, 45, 60)


def _fmt(value, decimals=0):
    if value is None or pd.isna(value):
        return "—"
    return f"{float(value):,.{decimals}f}"


def _money(value):
    if value is None or pd.isna(value):
        return "—"
    value = float(value)
    if abs(value) >= 1_000_000_000:
        return f"Rp {value / 1_000_000_000:,.2f} M"
    if abs(value) >= 1_000_000:
        return f"Rp {value / 1_000_000:,.1f} jt"
    return f"Rp {value:,.0f}"


def _replace_paragraphs(shape, texts, font_size=None):
    """Replace text while retaining existing paragraph positions and styling."""
    if not shape.has_text_frame:
        return
    paragraphs = list(shape.text_frame.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        text = texts[index] if index < len(texts) else ""
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ""
        elif text:
            paragraph.add_run().text = text
        if font_size is not None:
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
    for text in texts[len(paragraphs):]:
        paragraph = shape.text_frame.add_paragraph()
        paragraph.text = text
        if font_size is not None:
            for run in paragraph.runs:
                run.font.size = Pt(font_size)


def _remove_pictures(slide):
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shape._element.getparent().remove(shape._element)


def _add_table(slide, x, y, w, h, headers, rows, font_size=8.5):
    rows = list(rows)
    table_shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)
    )
    table = table_shape.table
    table.first_row = True
    for col_index in range(len(headers)):
        table.columns[col_index].width = Inches(w / len(headers))
    for row_index, values in enumerate([headers] + rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.margin_left = Inches(0.045)
            cell.margin_right = Inches(0.045)
            cell.margin_top = Inches(0.015)
            cell.margin_bottom = Inches(0.015)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if row_index == 0 else (LIGHT if row_index % 2 else WHITE)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT if col_index == 0 else PP_ALIGN.RIGHT
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size if row_index else font_size + 0.5)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = WHITE if row_index == 0 else DARK
    return table_shape


def _replace_chart(chart, categories, series):
    data = CategoryChartData()
    data.categories = [str(x) for x in categories]
    for name, values in series:
        data.add_series(str(name), [float(v or 0) for v in values])
    chart.replace_data(data)


def _safe_records(frame, columns):
    if frame is None or frame.empty:
        return [["—" for _ in columns]]
    return [[row.get(column, "—") for column in columns] for row in frame.to_dict("records")]


def generate_meeting_kasie_pptx(template_path, payload):
    """Fill the existing six-slide branch meeting template with monthly data."""
    template_path = Path(template_path)
    if not template_path.exists():
        raise FileNotFoundError(f"Template tidak ditemukan: {template_path}")
    prs = Presentation(str(template_path))
    if len(prs.slides) < 6:
        raise ValueError("Template Meeting Kasie harus memiliki enam slide.")

    month = payload["month_name"]
    year = int(payload["year"])
    business_label = payload["business_label"]
    period_label = f"{month} {year}"

    # Slide 1: cover.
    _replace_paragraphs(prs.slides[0].shapes[0], [
        f"MEETING KASIE {year}",
        f"MONTHLY BUSINESS REVIEW · {business_label.upper()}",
        "",
        "Cabang : BANDUNG",
        "",
        period_label,
    ])

    # Slide 3: current-year trend, principal contribution, KPI strips and customer summary.
    s3 = prs.slides[2]
    _replace_paragraphs(s3.shapes[2], [f"1. PERFORMANCE SALES · {period_label.upper()}"])
    _replace_paragraphs(s3.shapes[3], ["SUMMARY :"] + payload["sales_insight_lines"][:5], font_size=11)
    month_labels = payload["trend_month_labels"]
    monthly_actual = payload["trend_actual"]
    monthly_target = payload["trend_target"]
    average_2025 = payload["average_2025"]
    _replace_chart(s3.shapes[0].chart, month_labels, [
        ("ACTUAL SALES", monthly_actual), ("TARGET SALES", monthly_target),
        ("AVG 2025", [average_2025] * len(month_labels)),
    ])
    principal_names = payload["principal_contribution_names"]
    principal_values = payload["principal_contribution_values"]
    if not principal_names:
        principal_names, principal_values = ["Tidak ada data"], [0]
    _replace_chart(s3.shapes[1].chart, principal_names, [("ACTUAL SALES", principal_values)])
    _remove_pictures(s3)
    _add_table(s3, 0.1, 3.48, 13.05, 0.65,
               ["ACTUAL GROSS", "TARGET GROSS", "A/T", "DEVIASI", "POT"], [[
                   _money(payload["actual"]), _money(payload["target"]),
                   f"{payload['achievement']:.1f}%" if payload["achievement"] is not None else "—",
                   _money(payload["deviation"]), str(payload["pot_count"]),
               ]], font_size=9)
    _add_table(s3, 0.1, 4.25, 13.05, 0.65,
               ["ACTUAL YTD", "TARGET YTD", "A/T YTD", "DEV YTD", "PERIODE"], [[
                   _money(payload["ytd_actual"]), _money(payload["ytd_target"]),
                   f"{payload['ytd_achievement']:.1f}%" if payload["ytd_achievement"] is not None else "—",
                   _money(payload["ytd_deviation"]), f"Jan–{month}",
               ]], font_size=9)
    _add_table(s3, 0.1, 5.0, 6.25, 2.0,
               ["RELATION / CUSTOMER", "ACTUAL", "VS BULAN LALU"],
               _safe_records(payload["customer_table"], ["Customer", "Actual Text", "Change Text"]), font_size=8)

    # Slide 4: POT summary, detailed table and target-attainment chart.
    s4 = prs.slides[3]
    _replace_paragraphs(s4.shapes[0], [f"2. PRINCIPAL ON TARGET · {period_label.upper()}"])
    _replace_paragraphs(s4.shapes[1], ["POT · PENCAPAIAN ≥100%:"] + payload["pot_insight_lines"][:1], font_size=10)
    _replace_paragraphs(s4.shapes[2], ["PERLU PERHATIAN:"] + payload["below_insight_lines"][:2], font_size=10)
    _replace_paragraphs(s4.shapes[9], ["ANALISIS PERUBAHAN SALES:"] + payload["trend_insight_lines"][:4], font_size=8.5)
    _replace_paragraphs(s4.shapes[10], [f"POT POSITION · {period_label.upper()}"])
    _remove_pictures(s4)
    _add_table(s4, 0.08, 0.55, 6.9, 3.55,
               ["PRINCIPAL", "ACTUAL", "TARGET", "A/T", "DEV"],
               _safe_records(payload["principal_table"], ["Principal ID", "Actual Text", "Target Text", "AT Text", "Dev Text"]),
               font_size=7.4)
    chart_data = CategoryChartData()
    chart_data.categories = payload["attainment_chart_names"] or ["Tidak ada data"]
    chart_data.add_series("A/T %", payload["attainment_chart_values"] or [0])
    s4.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(7.1), Inches(0.48), Inches(6.0), Inches(3.52), chart_data)

    # Slide 5: monthly business group, program push, and customer movement.
    s5 = prs.slides[4]
    _replace_paragraphs(s5.shapes[0], [f"3. BUSINESS & PROGRAM REVIEW · {period_label.upper()}"])
    _replace_paragraphs(s5.shapes[1], payload["top_group_lines"][:2], font_size=11)
    _replace_paragraphs(s5.shapes[2], payload["principal_score_lines"][:2], font_size=11)
    _replace_paragraphs(s5.shapes[3], payload["meeting_summary_lines"][:6], font_size=9.5)
    _remove_pictures(s5)
    _add_table(s5, 0.12, 0.55, 9.35, 4.95,
               ["BISNIS GROUP", "ACTUAL", "TARGET", "A/T", "DEV"],
               _safe_records(payload["business_table"], ["Bisnis Group", "Actual Text", "Target Text", "AT Text", "Dev Text"]),
               font_size=8.3)
    _add_table(s5, 0.2, 5.66, 4.25, 1.55,
               ["PROGRAM PUSH", "TARGET", "ACTUAL", "A/T"],
               _safe_records(payload["program_table"], ["Program", "Target Text", "Actual Text", "AT Text"]),
               font_size=7.4)
    _add_table(s5, 4.58, 5.66, 8.55, 1.55,
               ["PERUBAHAN CUSTOMER", "ACTUAL", "VS BULAN LALU"],
               _safe_records(payload["customer_table"].head(4), ["Customer", "Actual Text", "Change Text"]),
               font_size=7.6)

    output = BytesIO()
    prs.save(output)
    return output.getvalue()
